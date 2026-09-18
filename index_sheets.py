#!/usr/bin/env python3
import os
import re
import sqlite3
import openpyxl
import xlrd
from datetime import datetime

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sheets_index.db")

def normalize_phone(val):
    if val is None:
        return None
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    s = str(val).strip()
    if s.endswith('.0') and s[:-2].replace('-', '').replace('+', '').isdigit():
        s = s[:-2]
    digits = re.sub(r'\D', '', s)
    if digits.startswith('20') and len(digits) in (12, 13, 14):
        # 2010..., 2011..., 2012..., 2015...
        if len(digits) == 12:
            return '0' + digits[2:]
    if len(digits) == 10 and digits[0] == '1': # 10..., 11..., 12..., 15...
        return '0' + digits
    if len(digits) == 11 and digits.startswith('01'):
        return digits
    return None

def normalize_arabic(text):
    if not text:
        return ""
    text = re.sub(r'[إأآا]', 'ا', text)
    text = re.sub(r'ة', 'ه', text)
    text = re.sub(r'ى', 'ي', text)
    return text.strip()

def init_db(conn):
    cur = conn.cursor()
    cur.execute("PRAGMA journal_mode = WAL;")
    cur.execute("PRAGMA synchronous = NORMAL;")
    
    # Files table
    cur.execute("""
    CREATE TABLE IF NOT EXISTS files (
        file_id INTEGER PRIMARY KEY AUTOINCREMENT,
        file_path TEXT UNIQUE,
        filename TEXT,
        folder TEXT,
        indexed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # Calls / CDR records table
    cur.execute("""
    CREATE TABLE IF NOT EXISTS cdr_records (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        file_id INTEGER,
        sheet_name TEXT,
        row_idx INTEGER,
        target_msisdn TEXT,
        target_norm TEXT,
        other_msisdn TEXT,
        other_norm TEXT,
        other_name TEXT,
        other_name_norm TEXT,
        event_time TEXT,
        duration TEXT,
        direction TEXT,
        other_id TEXT,
        other_address TEXT,
        cell_id TEXT,
        cell_address TEXT,
        raw_row TEXT,
        FOREIGN KEY(file_id) REFERENCES files(file_id)
    );
    """)

    # Indices for blazing fast search
    cur.execute("CREATE INDEX IF NOT EXISTS idx_cdr_target_norm ON cdr_records(target_norm);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_cdr_other_norm ON cdr_records(other_norm);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_cdr_other_name ON cdr_records(other_name_norm);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_cdr_other_id ON cdr_records(other_id);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_cdr_event_time ON cdr_records(event_time);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_cdr_file_sheet_row ON cdr_records(file_id, sheet_name, row_idx);")

    # Universal cell index (FTS or key-value) for arbitrary keyword lookups
    cur.execute("""
    CREATE VIRTUAL TABLE IF NOT EXISTS universal_search USING fts5(
        file_path UNINDEXED,
        sheet_name UNINDEXED,
        row_idx UNINDEXED,
        content,
        tokenize='unicode61'
    );
    """)
    conn.commit()

def parse_workbook(fpath):
    rows_data = []
    ext = os.path.splitext(fpath)[1].lower()
    try:
        if ext == '.xlsx':
            wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
            for sname in wb.sheetnames:
                sheet = wb[sname]
                for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                    if any(c is not None and str(c).strip() for c in row):
                        rows_data.append((sname, r_idx + 1, list(row)))
            wb.close()
        elif ext == '.xls':
            wb = xlrd.open_workbook(fpath, formatting_info=False)
            for sname in wb.sheet_names():
                sheet = wb.sheet_by_name(sname)
                for r_idx in range(sheet.nrows):
                    row = sheet.row_values(r_idx)
                    if any(c is not None and str(c).strip() for c in row):
                        rows_data.append((sname, r_idx + 1, row))
    except Exception as e:
        print(f"Error reading {fpath}: {e}")
    return rows_data

def process_file(fpath, conn):
    cur = conn.cursor()
    folder, filename = os.path.split(fpath)
    cur.execute("INSERT OR IGNORE INTO files (file_path, filename, folder) VALUES (?, ?, ?);", (fpath, filename, folder))
    cur.execute("SELECT file_id FROM files WHERE file_path = ?;", (fpath,))
    file_id = cur.fetchone()[0]

    # Delete previous entries if any
    cur.execute("DELETE FROM cdr_records WHERE file_id = ?;", (file_id,))
    cur.execute("DELETE FROM universal_search WHERE file_path = ?;", (fpath,))

    all_rows = parse_workbook(fpath)
    if not all_rows:
        return 0

    # Group by sheet
    sheets = {}
    for sname, r_idx, row in all_rows:
        sheets.setdefault(sname, []).append((r_idx, row))

    cdr_batch = []
    fts_batch = []

    for sname, sheet_rows in sheets.items():
        # Find headers
        header_map = {}
        for r_idx, row in sheet_rows[:20]:
            cleaned_row = [str(c).lower().strip() if c is not None else '' for c in row]
            if any('msisdn' in c or 'target' in c or 'call type' in c or 'called number' in c or 'sub_id' in c for c in cleaned_row):
                for idx, c in enumerate(cleaned_row):
                    if c:
                        header_map[c] = idx
                break

        for r_idx, row in sheet_rows:
            # Build text string for universal FTS
            row_non_empty = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if not row_non_empty:
                continue
            row_str = " | ".join(row_non_empty)
            fts_batch.append((fpath, sname, r_idx, row_str))

            # CDR extraction
            target_msisdn = None
            other_msisdn = None
            other_name = None
            event_time = None
            duration = None
            direction = None
            other_id = None
            other_address = None
            cell_id = None
            cell_address = None

            if header_map:
                def get_c(key):
                    for k, idx in header_map.items():
                        if key in k and idx < len(row):
                            return row[idx]
                    return None

                target_msisdn = get_c('target_msisdn') or get_c('msisdn') or get_c('a number')
                other_msisdn = get_c('other_msisdn') or get_c('called number') or get_c('b number') or get_c('dialed')
                other_name = get_c('other_name') or get_c('first_name')
                event_time = get_c('event_start_time') or get_c('call date') or get_c('full date')
                duration = get_c('call_duration') or get_c('duration') or get_c('volume')
                direction = get_c('event_direction') or get_c('call type')
                other_id = get_c('other_id') or get_c('sub_id_val')
                other_address = get_c('other_address') or get_c('street') or get_c('neighborhood')
                cell_id = get_c('cell_nid') or get_c('cell location')
                cell_address = get_c('cell_address') or get_c('district')
            else:
                # Heuristic mapping for standard 20-col telecom sheets
                if len(row) >= 9:
                    target_msisdn = row[1] if len(row) > 1 else None
                    event_time = row[5] if len(row) > 5 else None
                    direction = row[7] if len(row) > 7 else None
                    other_msisdn = row[8] if len(row) > 8 else None
                    other_name = row[9] if len(row) > 9 else None
                    other_id = row[10] if len(row) > 10 else None
                    other_address = row[12] if len(row) > 12 else None
                    cell_id = row[13] if len(row) > 13 else None
                    cell_address = row[14] if len(row) > 14 else None

            # Normalize values
            t_norm = normalize_phone(target_msisdn)
            o_norm = normalize_phone(other_msisdn)
            on_norm = normalize_arabic(str(other_name)) if other_name else None

            # Only record if there is some phone or name or id
            if t_norm or o_norm or other_name or other_id:
                cdr_batch.append((
                    file_id, sname, r_idx,
                    str(target_msisdn) if target_msisdn is not None else None,
                    t_norm,
                    str(other_msisdn) if other_msisdn is not None else None,
                    o_norm,
                    str(other_name) if other_name is not None else None,
                    on_norm,
                    str(event_time) if event_time is not None else None,
                    str(duration) if duration is not None else None,
                    str(direction) if direction is not None else None,
                    str(other_id) if other_id is not None else None,
                    str(other_address) if other_address is not None else None,
                    str(cell_id) if cell_id is not None else None,
                    str(cell_address) if cell_address is not None else None,
                    row_str
                ))

    if cdr_batch:
        cur.executemany("""
        INSERT INTO cdr_records (
            file_id, sheet_name, row_idx,
            target_msisdn, target_norm, other_msisdn, other_norm,
            other_name, other_name_norm, event_time, duration,
            direction, other_id, other_address, cell_id, cell_address, raw_row
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, cdr_batch)

    if fts_batch:
        cur.executemany("""
        INSERT INTO universal_search (file_path, sheet_name, row_idx, content)
        VALUES (?, ?, ?, ?);
        """, fts_batch)

    conn.commit()
    return len(cdr_batch)

def build_index():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    conn = sqlite3.connect(DB_PATH)
    init_db(conn)

    files_to_scan = []
    for root, dirs, files in os.walk(base_dir):
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            if ext in ['.xlsx', '.xls']:
                files_to_scan.append(os.path.join(root, f))

    print(f"Indexing {len(files_to_scan)} files into {DB_PATH}...")
    start_time = datetime.now()
    total_records = 0

    for i, fpath in enumerate(files_to_scan, 1):
        rel = os.path.relpath(fpath, base_dir)
        count = process_file(fpath, conn)
        total_records += count
        if i % 10 == 0 or i == len(files_to_scan):
            print(f"[{i}/{len(files_to_scan)}] Indexed {rel} ({count} records)")

    conn.close()
    duration = (datetime.now() - start_time).total_seconds()
    print(f"\nDone! Indexed {total_records} records from {len(files_to_scan)} files in {duration:.1f}s.")
    print(f"Database: {DB_PATH}")

if __name__ == "__main__":
    build_index()
